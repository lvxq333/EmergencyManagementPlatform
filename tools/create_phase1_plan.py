from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path("output/第一阶段详细开发规划与内容规格评审稿.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_font(run, name="Arial Unicode MS", size=10.5, bold=None, color="000000"):
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, before=0, after=5, line=1.22):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_para(doc, text="", style=None, bold_lead=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    format_paragraph(p)
    p.paragraph_format.keep_with_next = keep
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(item)
        set_font(r)
        format_paragraph(p, after=2)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_font(r)
        format_paragraph(p, after=2)


def add_table(doc, headers, rows, widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, size=font_size, bold=True, color="FFFFFF")
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
    for ri, row in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=font_size)
            format_paragraph(p, after=0, line=1.12)
            if ri % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    format_paragraph(p, before=8 if level == 1 else 5, after=5)
    p.paragraph_format.keep_with_next = True
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Bullet 2", "List Number"]:
    style = styles[style_name]
    style.font.name = "Arial Unicode MS"
    style_r_fonts = style._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        style_r_fonts.set(qn(f"w:{key}"), "Arial Unicode MS")
    style.font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True
styles["Heading 3"].font.size = Pt(11)
styles["Heading 3"].font.bold = True
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
remove_paragraph_borders(title)
run = title.add_run("课题五智能服务平台第一阶段详细开发规划与内容规格")
set_font(run, size=24, bold=True)
format_paragraph(title, before=18, after=8)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("范围和接口冻结阶段评审稿")
set_font(r, size=14, bold=True)
format_paragraph(subtitle, after=18)

meta_rows = [
    ("规划周期", "2026年9月5日至2026年9月13日"),
    ("版本", "V0.1 评审稿"),
    ("编制日期", "2026年9月5日"),
    ("适用对象", "项目负责人、平台开发、算法与数据团队、测试与第三方检测对接人员"),
    ("评审结论目标", "冻结第一阶段 P0 范围、数据标准、API 边界、模型协议、预警状态机和验收口径"),
]
add_table(doc, ["项目", "内容"], meta_rows, [1.35, 5.45], 9.5)

add_para(doc, "本稿用于第一阶段内部评审。阶段结束时不要求完成全部业务功能编码，但必须形成后续开发可直接执行、可测试、可追溯的统一基线。任何影响数据库主结构、接口字段、风险口径或预警流转的分歧，应在9月13日前形成书面结论。", bold_lead="本稿用途：")

add_heading(doc, "一 阶段结论", 1)
add_para(doc, "第一阶段的核心产出不是新增页面，而是把现有演示系统转化为可实施的工程合同。9月13日前应锁定“做什么、数据长什么样、接口如何调用、模型如何接入、预警如何流转、最终怎样验收”。完成该阶段后，第二阶段可直接进入设备台账、统一接入和300站仿真底座开发，不再围绕字段和流程反复返工。")
add_bullets(doc, [
    "范围冻结：形成任务书指标到平台模块、接口、数据表和测试证据的映射。",
    "数据冻结：统一设备、站点、观测值、数据集、风险结果、预警与处置记录的标识和字段。",
    "接口冻结：确定前端、接入层、模型服务和业务服务之间的 API V1 边界。",
    "流程冻结：确定预警生成、确认、派单、处置、解除和关闭的状态机及角色权限。",
    "验收冻结：与第三方检测方确认300站、2000并发、持续时间、错误率、响应时间和取证方式。",
])

add_heading(doc, "二 依据和当前基线", 1)
add_para(doc, "本规划依据《平台后续开发与结题实施规划》、当前前端页面、Spring Boot 后端、数据库脚本、Nginx 部署配置和现有接口文档编制。当前代码处于演示原型向正式平台过渡阶段。")
baseline_rows = [
    ("前端", "已具备综合态势、数据接入、短临降雨、燃气、供水、供热、耦合研判、预警处置、历史分析页面", "大量业务数据、规则和处置操作仍在浏览器内模拟"),
    ("后端", "Spring Boot 已有认证、用户、角色、权限、风险概览服务", "缺少设备、监测、数据集、模型结果、规则、事件闭环和历史聚合服务"),
    ("数据", "MySQL 已有用户权限、风险快照和简单告警表；Redis 用于风险状态", "缺少统一设备模型、观测数据、质量标识、数据集版本和来源追溯"),
    ("部署", "已具备 Nginx 静态代理和 Java 服务部署说明", "尚未形成结题环境配置基线、备份恢复和测试取证规范"),
    ("测试", "已有认证、用户、角色和风险服务自动化测试", "尚无统一数据契约测试、完整业务流测试和性能测试方案"),
]
add_table(doc, ["领域", "已有基础", "第一阶段需解决的设计缺口"], baseline_rows, [1.0, 2.85, 2.95], 8.7)

add_heading(doc, "三 阶段范围", 1)
add_heading(doc, "3.1 本阶段必须完成", 2)
scope_rows = [
    ("P1-S01", "任务书指标矩阵", "任务书目标、P0功能、平台模块、负责人、接口或数据对象、测试证据、完成节点"),
    ("P1-S02", "业务范围清单", "三类灾种、三类管线系统、三处示范区域占位、累计200公里管线口径"),
    ("P1-S03", "数据标准 V1", "站点设备、观测指标、数据集、质量码、风险结果、预警事件、审计字段"),
    ("P1-S04", "API 规范 V1", "认证、设备、数据接入、观测查询、场景回放、风险、预警、历史与系统接口"),
    ("P1-S05", "模型接口规范 V1", "统一请求响应、模型版本、输入数据集、证据项、置信度、错误码和超时"),
    ("P1-S06", "预警状态机 V1", "生成、确认、派单、处置、解除、关闭、重开及不可逆约束"),
    ("P1-S07", "样例数据包 V1", "CSV、JSON 模板；300站台账草案；正常、强降雨、燃气泄漏三个场景种子"),
    ("P1-S08", "测试口径与基线", "功能用例初稿、数据契约测试、性能口径、环境基线和问题清单"),
]
add_table(doc, ["编号", "交付项", "内容边界"], scope_rows, [0.75, 1.65, 4.4], 8.8)

add_heading(doc, "3.2 本阶段明确不做", 2)
add_bullets(doc, [
    "不在第一阶段完成300站持续上报、MQTT生产链路和正式时序存储，仅冻结方案并准备样例。",
    "不在第一阶段完成正式模型算法训练或准确率承诺，仅冻结模型服务接口和评测字段。",
    "不重构前端框架，不新增复杂三维 GIS、拖拽大屏、微服务拆分或可视化规则编排。",
    "不以仿真数据证明90%预警准确率；仿真只用于功能、流程、稳定性和性能验证。",
    "不在字段、状态机和测试口径未评审前大规模编写业务代码。",
])

add_heading(doc, "四 阶段出口和验收门槛", 1)
gate_rows = [
    ("G1 范围完整", "任务书 P0 指标100%进入追踪矩阵；每项均有负责人、实现载体和验证证据", "需求矩阵评审记录"),
    ("G2 数据可实现", "核心对象有唯一标识、字段类型、单位、枚举、必填性、质量码和时间语义", "数据字典 V1 与样例文件"),
    ("G3 接口可联调", "API 有路径、方法、认证、请求、响应、分页、错误码和幂等规则；OpenAPI 可解析", "OpenAPI V1 与契约检查"),
    ("G4 流程无歧义", "预警状态迁移、操作角色、必填处置内容、审计记录和异常分支均明确", "状态机评审与用例走查"),
    ("G5 样例可复现", "同一随机种子生成相同场景；样例可通过校验器；来源和数据集版本可追溯", "样例包与校验报告"),
    ("G6 验收口径明确", "300站、2000并发、持续时间、请求组合、统计方法、报告形式已书面确认或列明待确认责任人和期限", "测试口径确认单"),
    ("G7 开发可启动", "P0/P1优先级、第二阶段任务、依赖、风险和代码基线已冻结", "评审纪要与阶段签字"),
]
add_table(doc, ["出口门槛", "通过标准", "证据"], gate_rows, [1.25, 4.15, 1.4], 8.8)
add_para(doc, "判定规则：G1至G5和G7必须全部通过。G6如暂未取得第三方书面回复，可带条件通过，但必须记录联系人、待确认项、最晚回复日期和暂定内部控制值。", bold_lead="判定规则：")

add_heading(doc, "五 工作分解和内容指定", 1)
add_heading(doc, "5.1 需求追踪矩阵", 2)
trace_rows = [
    ("R01", "不少于300个站点或设备", "设备台账、接入监控", "station、device、latest_status", "台账导入、在线状态、持续上报测试"),
    ("R02", "至少两类接入方式", "统一数据接入", "ingest HTTP、file import、MQTT topic", "合法、缺字段、重复、乱序、异常单位测试"),
    ("R03", "气象甲烷压力位移等数据", "数据标准与历史", "observation、metric_definition、quality_code", "字段字典、单位换算、历史查询"),
    ("R04", "燃气供水供热三个系统", "三类管网专题", "asset、pipeline_segment、domain_risk", "每类系统典型场景和风险结果"),
    ("R05", "至少三类灾种或风险", "风险模型服务", "rain、leak、burst、heat、displacement 等枚举", "单灾种、耦合场景回放"),
    ("R06", "单系统与耦合风险", "风险计算和综合态势", "risk_result、risk_evidence、model_run", "来源、版本、贡献因子追溯"),
    ("R07", "精准预警和处置闭环", "预警处置", "rule、alert、work_order、action_log", "生成至关闭完整流程"),
    ("R08", "用户权限和审计", "系统管理", "user、role、permission、audit_log", "越权、禁用、锁定、关键操作审计"),
    ("R09", "2000人同时访问", "性能与部署", "测试脚本、监控指标、报告", "30分钟压测和错误率响应时间统计"),
    ("R10", "三处区域累计200公里", "示范应用", "region、pipeline_segment、evidence_attachment", "区域与里程统计及证明材料"),
]
add_table(doc, ["编号", "要求", "平台模块", "主要对象或接口", "验证证据"], trace_rows, [0.55, 1.45, 1.35, 1.65, 1.8], 8.0)

add_heading(doc, "5.2 业务范围和枚举", 2)
business_rows = [
    ("管线系统", "GAS 燃气、WATER 供水、HEAT 供热", "数据库、API、前端全部使用固定代码；中文仅作显示名"),
    ("首批风险类型", "RAIN 强降雨、GAS_LEAK 燃气泄漏、WATER_BURST 供水爆管", "第一阶段 P0；直接支撑首批三个场景"),
    ("扩展风险类型", "HEAT_ABNORMAL 热力异常、DISPLACEMENT 位移异常、COUPLING 多系统耦合", "接口中预留，第三阶段前实现"),
    ("运行模式", "SIMULATION 仿真、REPLAY 回放、REAL 真实", "所有结果必须携带 mode 和 datasetVersion"),
    ("风险等级", "LOW、MEDIUM、HIGH、CRITICAL", "建议阈值0–59、60–79、80–89、90–100；最终由评审冻结"),
    ("数据质量", "GOOD、SUSPECT、BAD、MISSING、LATE、DUPLICATE", "原始值不覆盖，转换值与质量原因并存"),
]
add_table(doc, ["分类", "固定内容", "使用约束"], business_rows, [1.25, 2.8, 2.75], 8.8)

add_heading(doc, "5.3 300站台账内容规格", 2)
station_rows = [
    ("GNSS及位移站", "100", "经纬度、高程、水平位移、垂直位移、沉降速率", "10秒至60秒"),
    ("气象监测站", "80", "温度、湿度、风速、风向、气压、瞬时雨量、累计雨量", "60秒至5分钟"),
    ("燃气监测终端", "50", "甲烷浓度、乙烷组分、压力、阴保电位", "5秒至30秒"),
    ("供水监测终端", "40", "压力、流量、流速、水锤标识", "5秒至30秒"),
    ("供热监测终端", "30", "供回水温度、压力、流量、补水量", "10秒至60秒"),
]
add_table(doc, ["站点类型", "数量", "核心指标", "建议频率"], station_rows, [1.55, 0.7, 3.4, 1.15], 8.8)
add_para(doc, "站点必填字段：stationId、stationName、stationType、regionCode、longitude、latitude、status、ownerOrg、protocol、reportIntervalSeconds、commissionedAt。设备必填字段：deviceId、stationId、deviceType、manufacturer、model、serialNumber、enabled、lastSeenAt。", bold_lead="站点必填字段：")
add_para(doc, "编号规则：站点采用 ST-{类型}-{四位序号}，例如 ST-GNSS-0001；设备采用 DV-{类型}-{五位序号}；管段采用 PL-{系统}-{区域}-{四位序号}。编号一旦发布不得复用。", bold_lead="编号规则：")

add_heading(doc, "5.4 观测数据标准", 2)
field_rows = [
    ("messageId", "string", "是", "单条消息唯一标识，用于幂等去重"),
    ("stationId, deviceId", "string", "是", "必须存在且处于启用状态"),
    ("metricCode", "enum", "是", "由指标字典定义，如 RAINFALL_5M、CH4_PPM、PRESSURE_MPA"),
    ("observedAt", "datetime", "是", "设备采集时间，ISO 8601，含时区"),
    ("receivedAt", "datetime", "系统生成", "平台接收时间，不允许客户端覆盖"),
    ("rawValue, rawUnit", "decimal, string", "是", "保留原始值和原始单位"),
    ("value, unit", "decimal, string", "是", "单位转换后的标准值和标准单位"),
    ("qualityCode", "enum", "是", "GOOD、SUSPECT、BAD、MISSING、LATE、DUPLICATE"),
    ("sourceType", "enum", "是", "HTTP、MQTT、FILE、SIMULATOR、REPLAY"),
    ("datasetId, scenarioId", "string", "条件必填", "仿真和回放必须填写；真实模式 datasetId 可按批次生成"),
    ("attributes", "object", "否", "厂商扩展字段，不进入核心风险计算前必须映射"),
]
add_table(doc, ["字段", "类型", "必填", "约束"], field_rows, [1.55, 1.0, 0.75, 3.5], 8.5)
add_bullets(doc, [
    "时间语义：统一存储 UTC 时间，接口按 ISO 8601 返回；前端默认显示 Asia/Shanghai。",
    "精度语义：经纬度保留至少6位小数；风险分数保留1位；测量值按指标字典规定精度。",
    "去重语义：优先以 messageId 去重；缺失时使用 deviceId、metricCode、observedAt 组合键。",
    "乱序语义：允许迟到数据入库但标记 LATE；实时状态不得被更早数据覆盖。",
    "质量语义：BAD 数据保存但默认不进入模型；SUSPECT 数据进入模型时必须在证据项中标记。",
])

add_heading(doc, "5.5 API V1 边界", 2)
api_rows = [
    ("认证权限", "POST /api/auth/login；GET /api/users；GET /api/roles；GET /api/permissions", "沿用现有接口，补充统一错误体和权限声明"),
    ("站点设备", "GET/POST /api/stations；GET/PUT /api/stations/{id}；GET /api/devices", "分页、区域、类型、在线状态筛选"),
    ("数据接入", "POST /api/ingest/observations；POST /api/import/jobs；GET /api/import/jobs/{id}", "HTTP 批量接入与文件导入任务；定义幂等键"),
    ("观测查询", "GET /api/observations/latest；GET /api/observations/history", "按站点、指标和时间范围查询"),
    ("场景数据", "GET/POST /api/scenarios；POST /api/scenarios/{id}/runs；POST /api/replays", "启动仿真或回放，返回 runId"),
    ("风险结果", "GET /api/risk/overview；GET /api/risk/results；GET /api/risk/results/{id}", "兼容现有概览，新增来源、版本和证据"),
    ("预警处置", "GET/POST /api/alert-rules；GET /api/alerts；POST /api/alerts/{id}/actions", "所有状态变化通过 action 接口并审计"),
    ("历史报表", "GET /api/analytics/trends；GET /api/analytics/regions；POST /api/reports", "第一阶段冻结查询参数和返回结构"),
    ("系统运维", "GET /api/health；GET /api/metrics/ingestion；GET /api/audit-logs", "健康、接入质量、审计查询"),
]
add_table(doc, ["域", "拟冻结接口", "第一阶段指定内容"], api_rows, [1.1, 3.65, 2.05], 8.1)
add_bullets(doc, [
    "认证：除登录、注册和健康检查外均默认要求 Bearer JWT；权限按 permissionKey 声明。",
    "响应：成功体统一包含 data、timestamp、requestId；列表包含 page、size、total。",
    "错误：统一包含 code、message、requestId、details；不得向前端暴露数据库异常堆栈。",
    "版本：当前采用 /api 路径并在 OpenAPI 的 info.version 标注 v1；破坏性变更进入 /api/v2。",
    "幂等：接入、场景启动、预警动作和报告生成必须支持 Idempotency-Key 或业务唯一键。",
])

add_heading(doc, "5.6 模型服务接口", 2)
model_rows = [
    ("请求身份", "requestId、modelCode、modelVersion、mode、datasetId、scenarioId、requestedAt"),
    ("空间范围", "regionCode、assetIds 或 geometryRef，不在业务请求中直接传超大 GeoJSON"),
    ("时间范围", "windowStart、windowEnd、forecastHorizonMinutes"),
    ("输入引用", "observationQuery 或 inputArtifactUri；记录输入数据版本和质量摘要"),
    ("结果主体", "resultId、riskType、riskScore、riskLevel、confidence、validFrom、validTo"),
    ("证据项", "metricCode、value、unit、threshold、contribution、sourceId、qualityCode"),
    ("可追溯信息", "algorithmName、modelVersion、parameterVersion、datasetVersion、executedAt、durationMs"),
    ("异常约束", "超时、输入不足、版本不存在、计算失败需使用固定错误码；不得返回虚构成功结果"),
]
add_table(doc, ["部分", "固定字段和约束"], model_rows, [1.35, 5.45], 8.8)
add_para(doc, "首批模型代码建议冻结为 NOWCAST_RAIN_V1、GAS_LEAK_RISK_V1、WATER_BURST_RISK_V1、HEAT_ABNORMAL_RISK_V1、DISPLACEMENT_RISK_V1、COUPLING_RISK_V1。第一阶段允许模型实现为可复现的规则或仿真适配器，但请求响应结构必须与正式模型一致。")

add_heading(doc, "5.7 预警状态机和角色权限", 2)
state_rows = [
    ("NEW 新生成", "规则引擎", "ACKNOWLEDGE、SUPPRESS", "需记录规则、风险结果、首次触发时间"),
    ("ACKNOWLEDGED 已确认", "值班员及以上", "ASSIGN、CLOSE_FALSE_POSITIVE", "确认意见必填"),
    ("ASSIGNED 已派单", "值班长或管理员", "START_PROCESS、REASSIGN", "责任人、截止时间必填"),
    ("PROCESSING 处置中", "责任人", "RESOLVE、ESCALATE", "过程记录和附件可追加"),
    ("RESOLVED 已解除", "责任人或系统", "CLOSE、REOPEN", "解除依据和恢复数据必填"),
    ("CLOSED 已关闭", "值班长或管理员", "REOPEN", "关闭结论必填；原记录不可修改"),
    ("SUPPRESSED 已抑制", "规则引擎或管理员", "REOPEN", "去重、维护或已知故障原因必填"),
]
add_table(doc, ["状态", "可执行角色", "允许动作", "审计要求"], state_rows, [1.55, 1.45, 1.75, 2.05], 8.4)
add_bullets(doc, [
    "每次动作生成不可变 actionLog，保存操作者、前后状态、意见、附件引用、时间和 requestId。",
    "相同对象、相同规则在抑制窗口内只更新计数和最高等级，不重复创建新预警。",
    "风险持续升高可升级；风险恢复只能进入 RESOLVED，不能由系统直接 CLOSED。",
    "关闭后再次触发默认创建关联新预警；只有人工判断为同一事件时才 REOPEN。",
])

add_heading(doc, "5.8 前端内容替换清单", 2)
ui_rows = [
    ("综合态势", "风险分数、区域分布、待处置预警、24小时趋势", "risk overview、alerts、analytics trends", "删除正式流程中的随机数；保留 demo 标识开关"),
    ("数据接入", "数据源、在线率、延迟、完整率、异常原因", "stations、devices、ingestion metrics", "明确站点数口径与质量码"),
    ("短临降雨", "实况、预测、区域排行、影响管网", "observations、model results", "展示模型版本、预测时效和数据来源"),
    ("三类管网", "核心指标、风险分布、贡献因素、演化链", "observations、domain risk results", "风险因子由 evidence 返回"),
    ("耦合研判", "交叉区、构成、放大系数、传播链", "coupling results、spatial relations", "交叉区需有稳定编号和关联管段"),
    ("预警处置", "列表、规则、时间线、动作", "alert rules、alerts、actions", "刷新或换账号后状态不丢失"),
    ("历史分析", "趋势、区域对比、预警统计、导出", "analytics、reports", "筛选项与 API 参数一致"),
]
add_table(doc, ["页面", "需正式化内容", "数据来源", "冻结要求"], ui_rows, [1.15, 2.25, 1.75, 1.65], 8.2)

add_heading(doc, "六 样例数据包规格", 1)
scenario_rows = [
    ("SCN-NORMAL-001", "正常运行", "全部指标在合理区间小幅波动", "不产生业务预警", "2026090501"),
    ("SCN-RAIN-001", "强降雨", "降雨增强，含水与位移滞后上升", "触发降雨风险并影响交叉区", "2026090502"),
    ("SCN-GAS-001", "燃气泄漏", "甲烷持续升高、压力下降", "生成燃气预警并可完成闭环", "2026090503"),
]
add_table(doc, ["场景编号", "名称", "时间轴内容", "预期结果", "固定种子"], scenario_rows, [1.35, 1.05, 2.25, 1.65, 0.85], 8.5)
add_para(doc, "每个场景目录必须包含 scenario.json、stations.csv、observations.csv 或 observations.jsonl、expected_risks.json、expected_alerts.json 和 README.md。README 需说明场景目的、起止时间、时间加速倍率、关键拐点、预期风险等级、允许误差和复现命令。")
add_bullets(doc, [
    "CSV 使用 UTF-8 编码，首行为英文固定字段名，禁止合并单元格和隐藏列。",
    "JSON 时间统一 ISO 8601；数值字段不得用带单位字符串代替。",
    "样例中所有站点、设备、管段和区域编号必须能在台账中找到。",
    "仿真和回放页面必须显示明显的模式标识，导出报告也必须保留该标识。",
])

add_heading(doc, "七 日程和责任分工", 1)
schedule_rows = [
    ("9月5日", "启动与基线", "确认第一阶段范围；盘点页面、接口、数据库和模拟数据；建立问题清单", "基线盘点 V1"),
    ("9月6日", "需求矩阵", "把任务书指标拆到模块、数据、接口、测试证据和负责人", "需求追踪矩阵 V1"),
    ("9月7日", "业务口径", "确定三类管线、首批风险、三处示范区域占位、管线里程统计方法", "范围清单与术语表"),
    ("9月8日", "数据标准", "冻结站点、设备、观测、数据集、质量码、风险结果核心字段", "数据字典 V1"),
    ("9月9日", "接口规范", "冻结 API 路径、认证、分页、错误、幂等；补齐 OpenAPI 框架", "API 规范 V1"),
    ("9月10日", "模型与预警", "冻结模型请求响应、证据项、预警状态机、规则和角色权限", "模型规范与状态机 V1"),
    ("9月11日", "样例与测试", "准备三场景样例；编写校验规则、契约用例和性能口径确认单", "样例包与测试初稿"),
    ("9月12日", "联合走查", "前端、后端、算法、测试逐页逐接口走查；关闭阻塞分歧", "走查问题关闭记录"),
    ("9月13日", "冻结评审", "通过出口门槛；签署评审纪要；拆解第二阶段任务", "冻结基线 V1.0"),
]
add_table(doc, ["日期", "主题", "主要任务", "当日产出"], schedule_rows, [0.9, 1.15, 3.65, 1.1], 8.3)

role_rows = [
    ("项目负责人", "范围、优先级、外部口径和最终冻结", "需求矩阵、评审纪要、第三方确认单"),
    ("后端负责人", "数据模型、API、幂等、状态机、审计", "数据字典、OpenAPI、状态机"),
    ("前端负责人", "页面字段映射、加载态、空态、错误态和权限显示", "页面接口映射表"),
    ("算法负责人", "模型输入输出、版本、证据项、评测边界", "模型接口规范与示例"),
    ("数据负责人", "站点台账、指标字典、样例场景、质量规则", "300站草案与样例包"),
    ("测试负责人", "测试用例、契约校验、性能口径、取证规范", "测试计划与用例初稿"),
    ("运维负责人", "环境、配置、日志、监控和版本基线", "环境清单与部署基线"),
]
add_table(doc, ["角色", "主要责任", "必须提交"], role_rows, [1.25, 3.35, 2.2], 8.7)

add_heading(doc, "八 测试内容指定", 1)
test_rows = [
    ("契约测试", "OpenAPI 可解析；样例请求响应通过 schema；错误码和必填字段一致", "所有拟冻结接口"),
    ("数据质量", "缺字段、类型错误、单位错误、越界、重复、乱序、迟到、未知设备", "每种异常至少1条用例"),
    ("权限测试", "未登录、过期令牌、禁用账号、越权角色、管理员操作", "关键写接口100%覆盖"),
    ("状态机测试", "合法迁移、非法跳转、并发操作、重复动作、关闭后重触发", "每个状态和动作均覆盖"),
    ("可追溯测试", "从预警追溯至风险结果、模型版本、数据集、原始观测和处置日志", "三类首批场景"),
    ("性能口径", "300站上报；2000并发30分钟；核心查询P95小于2秒；错误率小于1%", "先作为内部控制值，待第三方确认"),
]
add_table(doc, ["测试类别", "第一阶段需冻结的内容", "覆盖要求"], test_rows, [1.25, 4.15, 1.4], 8.7)
add_para(doc, "取证要求：每个验收用例指定测试编号、前置条件、输入数据集版本、执行步骤、预期结果、实际结果、截图或日志、执行人、执行时间和代码版本。性能报告必须记录机器配置、JVM 参数、数据库配置、接口比例、数据规模和统计方法。", bold_lead="取证要求：")

add_heading(doc, "九 风险和依赖", 1)
risk_rows = [
    ("真实数据未到位", "高", "按统一标准先完成仿真与回放；真实数据后续只替换接入适配器", "9月13日确认可获得字段和样本时间"),
    ("模型交付延迟", "高", "正式接口不变，先用版本化规则或仿真适配器", "算法负责人确认输入输出"),
    ("第三方口径不明", "高", "发出书面问题清单；内部采用更严格控制值", "9月6日前发出，9月13日跟踪"),
    ("示范区域材料不足", "中", "先保留 regionCode、管段里程和附件字段，不虚构证明材料", "9月13日确定责任人"),
    ("范围继续扩张", "高", "P0 冻结；新增项进入变更单并评估结题影响", "立即执行"),
    ("现有未提交改动较多", "中", "记录当前基线哈希、环境和未提交文件；冻结后再实施", "9月5日至9月6日"),
]
add_table(doc, ["风险", "级别", "措施", "本阶段动作"], risk_rows, [1.35, 0.65, 3.25, 1.55], 8.5)

add_heading(doc, "十 评审需确认的决策", 1)
decisions = [
    "D01 三类 P0 风险是否确定为强降雨、燃气泄漏、供水爆管；热力异常和位移异常作为扩展。",
    "D02 风险等级是否采用 LOW、MEDIUM、HIGH、CRITICAL 四级，以及对应分数阈值。",
    "D03 第一阶段接口是否采用本稿 API V1 分域；是否需要在路径中直接加入 /v1。",
    "D04 数据存储是否采用 MySQL 保存业务主数据和结果、Redis 保存最新状态；高频时序库选型在第二阶段验证后确认。",
    "D05 预警是否必须经 RESOLVED 后人工 CLOSED，系统不得直接关闭。",
    "D06 2000并发测试的用户行为比例、持续30分钟口径和核心接口 P95 小于2秒是否可作为内部基线。",
    "D07 三处示范区域的名称、责任单位、管线类型、里程及可提供的证明材料。",
    "D08 真实数据、模型和第三方检测对接的具体负责人及最晚交付日期。",
]
add_numbered(doc, decisions)

add_heading(doc, "十一 冻结与变更规则", 1)
add_bullets(doc, [
    "9月13日形成 V1.0 冻结包，至少包含需求矩阵、数据字典、OpenAPI、模型规范、状态机、样例包、测试初稿和评审纪要。",
    "冻结后新增字段必须向后兼容；删除、改名、类型变化、枚举含义变化均视为破坏性变更。",
    "破坏性变更须提交变更单，写明原因、影响页面、数据库迁移、兼容策略、测试影响和批准人。",
    "代码、样例数据、文档和部署配置必须使用同一版本标识；评审稿中的占位内容不得混入验收版本。",
    "第二阶段启动条件为出口门槛通过，不以日历日期自动视为完成。",
])

add_heading(doc, "十二 第一阶段完成检查表", 1)
check_rows = [
    ("□", "任务书指标均进入追踪矩阵，P0/P1优先级明确"),
    ("□", "三类管线、首批风险、运行模式和风险等级完成评审"),
    ("□", "站点、设备、观测、数据集、风险、预警、审计字段完成冻结"),
    ("□", "API V1 和 OpenAPI 文档完成联合走查并可解析"),
    ("□", "模型接口与证据项获得算法和平台双方确认"),
    ("□", "预警状态机、角色权限和异常分支通过用例走查"),
    ("□", "300站台账草案和三个首批场景样例可重复生成"),
    ("□", "数据质量、契约、权限、状态机和可追溯测试用例形成初稿"),
    ("□", "第三方检测问题清单已发送并形成跟踪记录"),
    ("□", "代码与环境基线、风险清单、第二阶段任务拆解已归档"),
    ("□", "评审纪要记录通过项、条件通过项、未决项、责任人和截止日期"),
]
add_table(doc, ["状态", "检查项"], check_rows, [0.65, 6.15], 9.0)

add_para(doc, "评审建议：本稿先用于范围评审。D01至D08确认后，将评审结论写入 V1.0 冻结版，并据此拆解第二阶段“数据接入与300站底座”的开发任务。", bold_lead="评审建议：")

footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("课题五智能服务平台第一阶段规划  V0.1 评审稿")
set_font(r, size=8, color="666666")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
